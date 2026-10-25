"""Balance sheet: python3 tools/make_balance_sheet.py <defs.json> (export Defs to JSON first), then recalc in LibreOffice."""
import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter as L

D = json.load(open(__import__('sys').argv[1]))
U = D['units']
order = {'crown': 0, 'brood': 1, 'clockwork': 2}
rows = sorted(U.items(), key=lambda kv: (order[kv[1]['faction']], kv[1].get('worker', False) is False, kv[1]['cost']))

wb = Workbook()
F = 'Arial'
blue = Font(name=F, color='0000FF')
black = Font(name=F)
bold = Font(name=F, bold=True)
hdr_fill = PatternFill('solid', fgColor='DDDDDD')
yellow = PatternFill('solid', fgColor='FFFF00')
fac_fill = {'crown': 'FCE9D6', 'brood': 'E5F2DC', 'clockwork': 'DCE8F5'}

# ---------------- Settings
st = wb.active
st.title = 'Settings'
st['A1'] = 'Balance sheet settings'; st['A1'].font = Font(name=F, bold=True, size=13)
st['A3'] = 'Second resource worth, in first-resource units (edit these)'; st['A3'].font = bold
st['A4'] = 'Faction'; st['B4'] = 'Second resource'; st['C4'] = 'Worth (x first resource)'
for c in 'ABC': st[c + '4'].font = bold; st[c + '4'].fill = hdr_fill
sec = [('crown', 'Powder (vs Gold)', 3), ('brood', 'Essence (vs Biomass)', 5), ('clockwork', 'Energy (vs Ore)', 3)]
for i, (f, n, v) in enumerate(sec):
    r = 5 + i
    st.cell(r, 1, f).font = black; st.cell(r, 2, n).font = black
    c = st.cell(r, 3, v); c.font = blue; c.fill = yellow
st['C5'].comment = Comment('Assumption: how many gold one powder is worth. Change it and every Value column recalculates.', 'Claude')
st['A9'] = 'Reference target armor (for "DPS vs armor" column)'; st['A9'].font = bold
st['C9'] = 2; st['C9'].font = blue; st['C9'].fill = yellow
st['A10'] = 'Reference attacker damage (for "Effective HP" column)'; st['A10'].font = bold
st['C10'] = 10; st['C10'].font = blue; st['C10'].fill = yellow
st['A12'] = 'How to read'; st['A12'].font = bold
notes = [
 'Blue numbers are inputs copied from scripts/defs.gd (game v0.7.18). Black numbers are formulas.',
 'Yellow cells are assumptions you can change; the Units sheet recalculates.',
 'Armor is flat: each hit loses the target\'s armor (minimum 1 damage). Acid strips 2 armor.',
 'Supply is per creature. Brood eggs hatch "Spawn" creatures, so an egg costs Supply x Spawn.',
 'Power = DPS x Effective HP: a rough "fight value" of one unit (Lanchester-style).',
 'Splash, cones, piercing, charges, spells and anti-air are not in DPS: see the Abilities column.',
]
for i, n in enumerate(notes):
    st.cell(13 + i, 1, n).font = black
st.column_dimensions['A'].width = 52; st.column_dimensions['B'].width = 24; st.column_dimensions['C'].width = 22

# ---------------- Units
ws = wb.create_sheet('Units')
heads = ['Faction', 'Unit', 'Worker', 'Air', 'Hits air', 'Tags', 'Cost', 'Cost 2', 'Spawn', 'Supply', 'Train time (s)',
         'HP', 'Armor', 'Speed', 'Range', 'Damage', 'Cooldown (s)', 'Splash',
         # formulas
         'Value per creature', 'Supply per egg', 'DPS', 'DPS vs armor', 'Effective HP', 'Power (DPS x EHP)',
         'Power per value', 'Power per supply', 'DPS per value x100', 'EHP per value', 'Value per supply', 'Abilities / notes']
for i, h in enumerate(heads, 1):
    c = ws.cell(1, i, h); c.font = bold; c.fill = hdr_fill; c.alignment = Alignment(wrap_text=True, vertical='top')
ws.row_dimensions[1].height = 45
r = 2
for k, u in rows:
    f = u['faction']
    vals = [f, u['name'], 'yes' if u.get('worker') else '', 'yes' if u.get('air') else '',
            'yes' if (u.get('anti_air') or u.get('air')) and not u.get('ground_only') else '', ', '.join(u.get('tags', [])),
            u['cost'], u.get('cost2', 0), u.get('spawn', 1), u['supply'], u['time'], u['hp'], u['armor'], u['speed'],
            u['range'], u['damage'], u['cooldown'], u.get('splash', 0)]
    for i, v in enumerate(vals, 1):
        c = ws.cell(r, i, v); c.font = blue if i >= 7 else black
        c.fill = PatternFill('solid', fgColor=fac_fill[f]) if i <= 2 else PatternFill()
    sec_row = {'crown': 5, 'brood': 6, 'clockwork': 7}[f]
    fm = {
      19: f'=(G{r}+H{r}*Settings!$C${sec_row})/I{r}',
      20: f'=J{r}*I{r}',
      21: f'=IF(Q{r}>0,P{r}/Q{r},0)',
      22: f'=IF(Q{r}>0,IF(P{r}>0,MAX(1,P{r}-Settings!$C$9),0)/Q{r},0)',
      23: f'=L{r}*Settings!$C$10/MAX(1,Settings!$C$10-M{r})',
      24: f'=U{r}*W{r}',
      25: f'=IF(S{r}>0,X{r}/S{r},0)',
      26: f'=IF(J{r}>0,X{r}/J{r},0)',
      27: f'=IF(S{r}>0,U{r}/S{r}*100,0)',
      28: f'=IF(S{r}>0,W{r}/S{r},0)',
      29: f'=IF(J{r}>0,S{r}/J{r},0)',
    }
    for col, fml in fm.items():
        c = ws.cell(r, col, fml); c.font = black
    ab = ws.cell(r, 30, u.get('about', '')); ab.font = black; ab.alignment = Alignment(wrap_text=True, vertical='top')
    r += 1
last = r - 1
fmt = {19: '0.0', 21: '0.0', 22: '0.0', 23: '0', 24: '#,##0', 25: '0.0', 26: '#,##0', 27: '0.00', 28: '0.00', 29: '0.0', 17: '0.0#', 11: '0'}
for col, nf in fmt.items():
    for rr in range(2, last + 1):
        ws.cell(rr, col).number_format = nf
widths = [10, 18, 7, 5, 7, 10, 7, 7, 7, 7, 8, 6, 6, 6, 6, 7, 9, 6, 9, 8, 7, 8, 9, 11, 9, 10, 9, 8, 8, 80]
for i, w in enumerate(widths, 1):
    ws.column_dimensions[L(i)].width = w
ws.freeze_panes = 'C2'
ws.auto_filter.ref = f'A1:{L(len(heads))}{last}'
# faction averages (army only)
r = last + 3
ws.cell(r, 1, 'Army averages per faction (workers excluded)').font = bold
r += 1
for i, h in enumerate(['Faction', 'Units', 'Avg value/creature', 'Avg DPS', 'Avg EHP', 'Avg power per value', 'Avg value per supply', 'Max supply'], 1):
    c = ws.cell(r, i, h); c.font = bold; c.fill = hdr_fill
caps = {'crown': D['factions']['crown']['max_supply'], 'brood': D['factions']['brood']['max_supply'], 'clockwork': D['factions']['clockwork']['max_supply']}
for f in ['crown', 'brood', 'clockwork']:
    r += 1
    rng = lambda col: f'{col}2:{col}{last}'
    cond = f'{rng("A")},"{f}",{rng("C")},"<>yes"'
    ws.cell(r, 1, f).font = black
    ws.cell(r, 2, f'=COUNTIFS({cond})').font = black
    for col, src in zip([3, 4, 5, 6, 7], ['S', 'U', 'W', 'Y', 'AC']):
        c = ws.cell(r, col, f'=AVERAGEIFS({rng(src)},{cond})'); c.font = black; c.number_format = '0.0'
    c = ws.cell(r, 8, caps[f]); c.font = blue
r += 2
ws.cell(r, 1, 'Max army value at cap = Avg value per supply x (Max supply - workers). Clockwork: few, tough; Brood: many, cheap.').font = black

# ---------------- Buildings
bs = wb.create_sheet('Buildings')
bh = ['Faction', 'Building', 'Cost', 'Build time (s)', 'HP', 'Armor', 'Size', 'Supply', 'Crew', 'Per hall', 'Trains', 'Notes']
for i, h in enumerate(bh, 1):
    c = bs.cell(1, i, h); c.font = bold; c.fill = hdr_fill
r = 2
for k, b in sorted(D['buildings'].items(), key=lambda kv: (order[kv[1]['faction']], kv[1]['cost'])):
    f = b['faction']
    econ = D['factions'][f]['econ'].get(k, '')
    size = b['size']
    vals = [f, b['name'], b['cost'], b['time'], b['hp'], b['armor'], str(size).replace('(', '').replace(')', '').replace(', ', 'x'),
            b['supply'], b.get('crew', ''), econ, ', '.join(D['units'][x]['name'] for x in b['trains']), b.get('about', '')]
    for i, v in enumerate(vals, 1):
        c = bs.cell(r, i, v); c.font = blue if 3 <= i <= 10 else black
    bs.cell(r, 12).alignment = Alignment(wrap_text=True, vertical='top')
    r += 1
for i, w in enumerate([10, 20, 7, 9, 7, 6, 7, 7, 6, 8, 40, 80], 1):
    bs.column_dimensions[L(i)].width = w
bs.freeze_panes = 'C2'

out = 'docs/Holdfast_balance.xlsx'
wb.save(out)
print(out, last)
